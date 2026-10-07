# -*- coding: utf-8 -*-
"""识别结果写入 Excel 工作簿（openpyxl）。"""
from __future__ import annotations

import re
from typing import Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# 疑似数字的文本（含千分位、百分号、负号）转数值，便于在 Excel 里校验求和
_NUM_RE = re.compile(r"^-?\d{1,3}(,\d{3})*(\.\d+)?$")
_NUM_PLAIN_RE = re.compile(r"^-?\d+(\.\d+)?$")
_PCT_RE = re.compile(r"^(-?\d+(?:\.\d+)?)%$")

_HEADER_FILL = PatternFill("solid", fgColor="EAF2E8")
_HEADER_FONT = Font(bold=True)
_THIN = Side(style="thin", color="B0B0B0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def coerce_number(text: str):
    """纯数字文本转 int/float，其余原样返回（百分数按数值写入）。"""
    t = (text or "").strip()
    if not t:
        return t
    m = _PCT_RE.match(t)
    if m:
        return round(float(m.group(1)) / 100.0, 6)
    if _NUM_PLAIN_RE.match(t):
        return int(t) if re.match(r"^-?\d+$", t) else float(t)
    if _NUM_RE.match(t):
        return int(t.replace(",", "")) if "." not in t else float(t.replace(",", ""))
    return t


def safe_sheet_name(name: str, used: Dict[str, int]) -> str:
    """规范化 sheet 名（Excel 限制 31 字符与部分字符），冲突自动加序号。"""
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", (name or "Sheet")).strip() or "Sheet"
    cleaned = cleaned[:31]
    if cleaned not in used:
        used[cleaned] = 1
        return cleaned
    used[cleaned] += 1
    suffix = f"({used[cleaned]})"
    return (cleaned[:31 - len(suffix)] + suffix)


def write_workbook(out_path: str, pages: List[Dict]) -> None:
    """写入一个工作簿。

    pages: [{"name": sheet 名, "title": 表格标题(可空),
             "rows": [[str]], "merges": [(r,c,rspan,cspan)]}, ...]
    行列下标从 0 开始；merges 为空列表表示无合并单元格。
    有 title 时写为第 1 行（跨列合并居中），表头从第 2 行开始。
    """
    wb = Workbook()
    wb.remove(wb.active)
    used: Dict[str, int] = {}
    for page in pages:
        rows: List[List[str]] = page.get("rows") or []
        merges: List[Tuple[int, int, int, int]] = page.get("merges") or []
        title = (page.get("title") or "").strip()
        plain = bool(page.get("plain"))   # 整页文字模式：无标题、无表头样式
        ws = wb.create_sheet(title=safe_sheet_name(page.get("name") or "Sheet", used))
        offset = 1 if title else 0

        if title:
            n_cols = max((len(r) for r in rows), default=1)
            ws.cell(row=1, column=1, value=title)
            ws.cell(row=1, column=1).font = Font(bold=True, size=14)
            ws.cell(row=1, column=1).alignment = Alignment(
                horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 30
            if n_cols > 1:
                ws.merge_cells(start_row=1, start_column=1,
                               end_row=1, end_column=n_cols)

        for r, row in enumerate(rows, start=1 + offset):
            for c, text in enumerate(row, start=1):
                value = coerce_number(str(text))
                ws.cell(row=r, column=c, value=value)
                ws.cell(row=r, column=c).border = _BORDER
                ws.cell(row=r, column=c).alignment = Alignment(
                    vertical="center",
                    horizontal="left" if isinstance(value, str) else "right",
                )

        for (r, c, rspan, cspan) in merges:
            if rspan > 1 or cspan > 1:
                ws.merge_cells(start_row=r + 1 + offset, start_column=c + 1,
                               end_row=r + rspan + offset, end_column=c + cspan)

        _style_header(ws, rows, offset) if not plain else _style_plain(ws)
        _fit_columns(ws, rows, offset) if not plain else _fit_plain_columns(ws)

    if len(wb.sheetnames) == 0:  # 空数据也产出合法文件
        wb.create_sheet(title="Sheet")
    wb.save(out_path)


def _style_header(ws, rows: List[List[str]], offset: int = 0) -> None:
    """表头行（rows 第一行，含标题时位于工作表第 2 行）加底色加粗居中。"""
    if not rows:
        return
    header_row = 1 + offset
    n_cols = max(len(r) for r in rows)
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _fit_columns(ws, rows: List[List[str]], offset: int = 0) -> None:
    """按内容粗略自适应列宽（中文按 2 个字符宽度计），标题行不参与统计。"""
    if not rows:
        return
    n_cols = max(len(r) for r in rows)
    for c in range(1, n_cols + 1):
        width = 0
        for row in rows[:200]:
            text = str(row[c - 1]) if c <= len(row) else ""
            w = sum(2 if ord(ch) > 127 else 1 for ch in text)
            width = max(width, w)
        ws.column_dimensions[get_column_letter(c)].width = min(max(width + 3, 8), 42)


def _style_plain(ws) -> None:
    """整页文字模式：左对齐即可，首行不做表头强调。"""
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 500)):
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="center")


def _fit_plain_columns(ws) -> None:
    """整页文字模式：单列宽版面。"""
    if ws.max_column >= 1:
        ws.column_dimensions["A"].width = 90
