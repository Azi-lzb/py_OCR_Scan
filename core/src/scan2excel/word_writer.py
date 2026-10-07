# -*- coding: utf-8 -*-
"""识别结果导出 Word 文档（python-docx）。

每个识别页一个章节（分页符隔开）：表格页渲染为 Word 表格（含合并
单元格还原），文字页渲染为逐行段落。
"""
from __future__ import annotations

from typing import Dict, List, Tuple

from docx import Document
from docx.shared import Pt


def write_document(out_path: str, pages: List[Dict]) -> None:
    """写入一个 Word 文档。

    pages: [{"name": 页名, "title": 表格标题(可空), "mode"/"plain": 文字页标记,
             "rows": [[str]], "merges": [(r,c,rspan,cspan)]}, ...]
    """
    doc = Document()
    for i, page in enumerate(pages):
        if i:
            doc.add_page_break()
        rows: List[List[str]] = page.get("rows") or []
        merges: List[Tuple[int, int, int, int]] = page.get("merges") or []
        plain = bool(page.get("plain")) or page.get("mode") == "text"
        title = (page.get("title") or "").strip()
        heading = title if (title and not plain) else str(page.get("name") or f"第{i + 1}页")
        doc.add_heading(heading, level=1)

        if not rows:
            continue
        if plain:
            for row in rows:
                text = str(row[0]) if row else ""
                para = doc.add_paragraph(text)
                para.paragraph_format.space_after = Pt(4)
        else:
            n_cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=n_cols)
            table.style = "Table Grid"
            # 先填非覆盖格，再做合并（合并会拼接内容，覆盖格必须留空）
            covered = set()
            for (r, c, rspan, cspan) in merges:
                if rspan > 1 or cspan > 1:
                    for rr in range(r, r + rspan):
                        for cc in range(c, c + cspan):
                            if not (rr == r and cc == c):
                                covered.add((rr, cc))
            for r, row in enumerate(rows):
                for c in range(n_cols):
                    if (r, c) in covered:
                        continue
                    table.cell(r, c).text = str(row[c]) if c < len(row) else ""
            for (r, c, rspan, cspan) in merges:
                if rspan > 1 or cspan > 1:
                    table.cell(r, c).merge(table.cell(min(r + rspan - 1, len(rows) - 1),
                                                      min(c + cspan - 1, n_cols - 1)))
    doc.save(out_path)
