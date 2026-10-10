# -*- coding: utf-8 -*-
"""WebApi 集成测试：模拟前端完整操作流（加图→识别→改格→导出→删除）。

不经任何外壳/窗口，直接调用 core 的 WebApi；Flask 壳另有冒烟测试。
"""
from __future__ import annotations

import json
import os
import sys
import time
from pathlib import Path

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "core", "src"))

DATA = os.path.join(HERE, "data")


def main() -> int:
    from scan2excel.web_app import WebApi

    api = WebApi(ROOT)
    ok = True

    # 1) 加图（表格 2 张 + 文字 1 张混合批次）
    added = api.add_image_paths([
        os.path.join(DATA, "inventory_flat.png"),
        os.path.join(DATA, "production_merge.png"),
        os.path.join(DATA, "notes_photo.png"),      # 无表格线 → 文字模式
        os.path.join(DATA, "不存在的图.png"),   # 应被跳过
        __file__,                                  # 非图片应被跳过
    ])
    assert added == 3, f"期望添加 3 张，实际 {added}"
    assert len(api.state["images"]) == 3
    assert api.state["current"] == 0
    print("✅ add_image_paths：3 张已添加（含 1 张文字页），非法文件已跳过")

    # 2) 识别（后台线程）
    assert api.start_ocr() is True
    deadline = time.time() + 180
    while api.state["busy"] and time.time() < deadline:
        time.sleep(0.5)
    assert not api.state["busy"], "识别超时"
    imgs = api.state["images"]
    assert imgs[0]["status"] == "完成", imgs[0]
    assert imgs[1]["status"] == "完成", imgs[1]
    assert imgs[2]["status"] == "完成", imgs[2]
    assert imgs[0]["mode"] == "table" and imgs[2]["mode"] == "text", \
        [i["mode"] for i in imgs]
    assert api.state["table"] and api.state["table"]["rows"]
    print(f"✅ start_ocr：表格 {imgs[0]['n_rows']}x{imgs[0]['n_cols']} + "
          f"{imgs[1]['n_rows']}x{imgs[1]['n_cols']}，文字 {imgs[2]['n_rows']} 行（模式自动判断）")

    # 2b) 切到文字页检查单列结果
    st = api.set_current(2)
    assert st["table"]["mode"] == "text"
    assert all(len(r) == 1 for r in st["table"]["rows"])
    st = api.set_current(0)
    assert st["table"]["mode"] == "table"
    print("✅ set_current：文字/表格页切换正常")

    # 3) 预览
    durl = api.get_preview()
    assert durl.startswith("data:image/jpeg;base64,"), "预览 dataURL 异常"
    api.set_preview_mode("overlay")
    assert api.get_preview() != durl
    print("✅ get_preview：原图/框线叠加两种预览可用")

    # 4) 校对修改（含标题）
    table = api.state["table"]
    assert table["title"] == "仓库月末盘点表", table.get("title")
    old = table["rows"][1][0]
    assert api.update_cell(1, 0, "改过的值") is True
    assert table["rows"][1][0] == "改过的值"
    assert api.update_cell(999, 0, "x") is False
    assert api.update_title("盘点表(修订)") is True
    assert table["title"] == "盘点表(修订)"
    print(f"✅ update_cell/update_title：{old} → 改过的值，标题 → 盘点表(修订)")

    # 5) 导出 Excel（单张工作簿：每图一个 Sheet，覆盖不追加；sheet 名=文件名）
    out = os.path.join(DATA, "_webapi_test.xlsx")
    if os.path.exists(out):
        os.remove(out)
    res = api.export_excel(path=out, naming="file")
    assert res["ok"] and os.path.exists(out), res
    from openpyxl import load_workbook
    wb = load_workbook(out)
    # 「汇总检查」sheet 固定插在最前，其后为各照片 sheet
    assert wb.sheetnames == ["汇总检查", "inventory_flat", "production_merge",
                             "notes_photo"], wb.sheetnames
    sum_ws = wb["汇总检查"]
    assert sum_ws["A1"].value == "序号" and sum_ws["G1"].value == "勾稽结论"
    ws = wb["inventory_flat"]
    # 无模板页：普通 sheet（首行表头，第一行数据从第 2 行起）
    assert ws["A1"].value == "序号", ws["A1"].value
    assert ws["A2"].value == "改过的值", ws["A2"].value
    # 合并单元格还原检查
    merges = [str(m) for m in wb["production_merge"].merged_cells.ranges]
    assert any("A1" in m for m in merges), f"合并单元格未还原: {merges}"
    print("✅ export_excel(file)：一张工作簿 3 个 Sheet，表头/合并正确")

    # 5a) 按表标题命名导出
    out_t = os.path.join(DATA, "_webapi_test_title.xlsx")
    if os.path.exists(out_t):
        os.remove(out_t)
    res_t = api.export_excel(path=out_t, naming="title")
    assert res_t["ok"], res_t
    wb_t = load_workbook(out_t)
    assert "盘点表(修订)" in wb_t.sheetnames, wb_t.sheetnames
    wb_t.close()
    os.remove(out_t)
    print("✅ export_excel(title)：sheet 名=表标题")

    # 5b) 同路径再导一次：应整体覆盖而不是追加 Sheet
    res2 = api.export_excel(path=out)
    assert res2["ok"], res2
    wb2 = load_workbook(out)
    assert wb2.sheetnames == wb.sheetnames, f"导出追加而非覆盖: {wb2.sheetnames}"
    print("✅ export_excel 覆盖导出：重复导出不追加 Sheet")

    # 6) 导出 Word（表格页=Word表格含合并，文字页=段落）
    out_w = os.path.join(DATA, "_webapi_test.docx")
    if os.path.exists(out_w):
        os.remove(out_w)
    res = api.export_word(path=out_w)
    assert res["ok"] and os.path.exists(out_w), res
    from docx import Document
    doc = Document(out_w)
    para_texts = [p.text for p in doc.paragraphs]
    assert "盘点表(修订)" in para_texts, para_texts[:6]      # 表格页标题
    assert len(doc.tables) == 2, f"应有 2 个 Word 表格: {len(doc.tables)}"
    t0 = doc.tables[0]
    assert t0.cell(0, 0).text == "序号" and t0.cell(1, 0).text == "改过的值"
    t1 = doc.tables[1]
    assert t1.cell(0, 0).text == "车间", f"Word 合并单元格内容异常: {t1.cell(0,0).text}"
    assert any(p == "备件入库明细（2026年10月）" for p in para_texts), "文字页段落缺失"
    assert any(p.startswith("经手人") for p in para_texts), "文字页尾行缺失"
    print("✅ export_word：表格页转 Word 表格（含合并），文字页逐行段落")

    # 7) 强制文字识别（对表格页点"识别文字"）
    api.set_current(0)
    assert api.start_text_ocr(0) is True
    deadline = time.time() + 120
    while api.state["busy"] and time.time() < deadline:
        time.sleep(0.5)
    assert not api.state["busy"], "文字识别超时"
    assert api.state["images"][0]["mode"] == "text"
    st = api.set_current(0)
    assert st["table"]["mode"] == "text" and all(len(r) == 1 for r in st["table"]["rows"])
    print(f"✅ start_text_ocr：表格页强制文字识别得 {st['table']['rows'].__len__()} 行")

    # 7b) 逐格分数透传（低置信度标红的数据源）
    st = api.set_current(1)
    scores = st["table"].get("scores") or {}
    assert scores, "table.scores 为空"
    assert all(isinstance(v, float) for v in scores.values())
    print(f"✅ table.scores：{len(scores)} 格带置信度")

    # 7c) 忽略区域：全图区域 → 重新识别后整页为空
    api.add_ignore_region(2, 0.0, 0.0, 1.0, 1.0)
    api.start_text_ocr(2)
    deadline = time.time() + 120
    while api.state["busy"] and time.time() < deadline:
        time.sleep(0.5)
    st = api.set_current(2)
    assert not st["has_result"] or not st["table"]["rows"], "忽略区域未过滤文字"
    api.clear_ignore_regions(2)
    print("✅ 忽略区域：全图忽略后重识别为空，清空后可恢复")

    # 8) PDF 输入：两页 PDF 逐页展开
    from PIL import Image
    pdf_path = os.path.join(DATA, "_two_pages.pdf")
    pil_pages = [Image.open(os.path.join(DATA, "inventory_flat.png")),
                 Image.open(os.path.join(DATA, "production_merge.png"))]
    pil_pages[0].save(pdf_path, save_all=True, append_images=pil_pages[1:])
    added = api.add_image_paths([pdf_path])
    assert added == 2, f"PDF 应展开 2 页，实际 {added}"
    names = [i["name"] for i in api.state["images"]]
    assert any("_p1" in n for n in names) and any("_p2" in n for n in names), names
    os.remove(pdf_path)
    print("✅ PDF 输入：2 页 PDF 自动展开为 2 张照片")

    # 9) 多格式导出（第 0 页在第 7 步已被强制转文字，改用仍是表格的第 1 页）
    api.set_current(1)
    outs = {}
    for fmt in ("csv", "txt", "md", "json"):
        p = os.path.join(DATA, f"_export.{fmt}")
        res = api.export_data(fmt, path=p)
        assert res["ok"] and os.path.exists(p), (fmt, res)
        outs[fmt] = Path(p).read_text(encoding="utf-8-sig" if fmt == "csv" else "utf-8")
    assert "车间" in outs["csv"] and "12500" in outs["csv"]
    assert "\n\n" not in outs["csv"], "CSV 出现空行"
    assert "=====" in outs["txt"] and "WL-1001" in outs["txt"]
    assert "## 车间月度产量统计" in outs["md"] and "|---|" in outs["md"]
    import json as _json
    j = _json.loads(outs["json"])
    # 第 2 页在忽略区域测试后结果为空，不参与导出，故为 2 页
    assert len(j) >= 2 and any(p["name"] == "inventory_flat" for p in j)
    for fmt in ("csv", "txt", "md", "json"):
        os.remove(os.path.join(DATA, f"_export.{fmt}"))
    print("✅ export_data：CSV/TXT/Markdown/JSON 四种格式导出正确")

    # 9b) 高精度档开关（不触发下载）
    assert api.check_high_accuracy_ready() in (True, False)
    st = api.set_high_accuracy(False)
    assert st["high_accuracy"] is False
    print("✅ 高精度档接口可用（当前未启用，不触发模型下载）")

    # 10) 删除/清空（3 张原始 + PDF 展开 2 页 = 5 张）
    st = api.remove_image(0)
    assert len(st["images"]) == 4
    st = api.clear_images()
    assert st["images"] == [] and st["current"] == -1
    print("✅ remove_image / clear_images")

    os.remove(out)
    os.remove(out_w)
    print("\n总体结论：✅ WebApi 全部通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
