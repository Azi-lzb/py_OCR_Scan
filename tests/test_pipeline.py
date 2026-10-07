# -*- coding: utf-8 -*-
"""全链路测试（不经 UI）：测试图 → 切格 → OCR → Excel → 读回验证。

用法：
  .venv/Scripts/python.exe tests/test_pipeline.py            # 先自动生成测试图
"""
from __future__ import annotations

import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "core", "src"))

OUT_DIR = os.path.join(HERE, "data")
XLSX_PATH = os.path.join(OUT_DIR, "识别结果_测试.xlsx")


_PUNC_EQ = str.maketrans({
    "，": ",", "：": ":", "；": ";", "！": "!", "？": "?",
    "（": "(", "）": ")", "　": " ",
})


def norm(text: str) -> str:
    """宽松比对：忽略空白与全半角标点差异。"""
    t = (text or "").translate(_PUNC_EQ)
    return t.replace(" ", "").strip()


def compare(result_rows, expected_rows):
    """返回 (匹配格数, 非空总格数, 不匹配明细)。真值网格可能与识别网格尺寸不同时按位置比对。"""
    matched, total, diffs = 0, 0, []
    for r, erow in enumerate(expected_rows):
        for c, etext in enumerate(erow):
            if not norm(etext):
                continue
            total += 1
            got = result_rows[r][c] if r < len(result_rows) and c < len(result_rows[r]) else ""
            if norm(got) == norm(etext):
                matched += 1
            else:
                diffs.append((r, c, etext, got))
    return matched, total, diffs


def main() -> int:
    if not os.path.exists(os.path.join(OUT_DIR, "expected.json")):
        print("生成测试图 ...")
        rc = os.system(f'"{sys.executable}" "{os.path.join(HERE, "make_test_images.py")}"')
        if rc != 0:
            print("生成测试图失败")
            return 1

    from scan2excel.service import Scan2ExcelService
    from scan2excel.excel_writer import write_workbook

    with open(os.path.join(OUT_DIR, "expected.json"), encoding="utf-8") as f:
        expected = json.load(f)

    service = Scan2ExcelService()
    pages, all_ok = [], True
    for fname in sorted(expected.keys()):
        path = os.path.join(OUT_DIR, fname)
        print(f"\n=== {fname} ===")
        page = service.process_image(path, on_step=lambda s: print("  [step]", s))
        pages.append({"name": page.name, "title": page.title,
                      "plain": page.mode == "text",
                      "rows": page.rows, "merges": page.merges})
        if page.error:
            print(f"  ❌ 失败: {page.error}")
            all_ok = False
            continue
        exp = expected[fname]
        matched, total, diffs = compare(page.rows, exp["rows"])
        acc = matched / total * 100 if total else 0
        if exp.get("mode") == "text":
            # 文字模式：无标题概念，行文本匹配即可
            status = "✅" if acc >= 90 else "❌"
            print(f"  {status} 文字模式 {page.n_rows} 行, 耗时{page.elapsed:.1f}s, "
                  f"行文本匹配 {matched}/{total} = {acc:.1f}%")
            if acc < 90:
                all_ok = False
            continue
        if exp.get("borderless"):
            # 无框线表格：结构模型行列对齐可能与真值略有出入，放宽到 75%
            ok = page.mode == "table" and page.borderless and acc >= 75 and page.n_cols >= 3
            status = "✅" if ok else "❌"
            print(f"  {status} 无框线表格 {page.n_rows}行x{page.n_cols}列"
                  f"（{'走模型结构识别' if page.borderless else '❌未走结构识别'}）, "
                  f"耗时{page.elapsed:.1f}s, 匹配 {matched}/{total} = {acc:.1f}%")
            for (r, c, etext, got) in diffs[:6]:
                print(f"     · R{r}C{c}: 期望[{etext}] 实得[{got}]")
            if not ok:
                all_ok = False
            continue
        title_ok = norm(page.title) == norm(exp.get("title", ""))
        if exp.get("rotated"):
            # 横放照片：自动转正后应与平拍完全一致（含标题）
            ok = acc >= 95 and title_ok
            status = "✅" if ok else "❌"
            print(f"  {status} 横放照片 {page.n_rows}行x{page.n_cols}列, "
                  f"耗时{page.elapsed:.1f}s, 准确率 {acc:.1f}%, 标题[{page.title}]")
            if not ok:
                all_ok = False
            continue
        status = "✅" if (acc >= 90 and title_ok) else "❌"
        print(f"  {status} {page.n_rows}行x{page.n_cols}列, 耗时{page.elapsed:.1f}s, "
              f"识别准确率 {matched}/{total} = {acc:.1f}%, 最低置信度 {page.min_score}"
              + (f", 标题[{page.title}]" if page.title else ", 无标题"))
        if not title_ok and exp.get("title"):
            print(f"     · 标题: 期望[{exp['title']}] 实得[{page.title}]")
        for (r, c, etext, got) in diffs[:8]:
            print(f"     · R{r}C{c}: 期望[{etext}] 实得[{got}]")
        if acc < 90 or not title_ok:
            all_ok = False

    write_workbook(XLSX_PATH, pages)
    print(f"\n已导出: {XLSX_PATH}")

    # 读回验证
    from openpyxl import load_workbook
    wb = load_workbook(XLSX_PATH)
    print("sheet:", wb.sheetnames)
    ok_sheets = len(wb.sheetnames) == len(pages)
    ws = wb[wb.sheetnames[0]]
    assert ws.max_row >= 2 and ws.max_column >= 2, "工作表内容为空"
    print(f"首个 sheet 维度: {ws.max_row}行 x {ws.max_column}列")
    print("\n总体结论:", "✅ 全部通过" if (all_ok and ok_sheets) else "❌ 存在未达标项")
    return 0 if (all_ok and ok_sheets) else 2


if __name__ == "__main__":
    sys.exit(main())
